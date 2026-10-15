"""Build the three-page user handout from local RWS assets; no network access."""
from pathlib import Path
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "DocumentCompare-QuickGuide.docx"
PURPLE, PINK = "3E016F", "E60054"
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Mm(210), Mm(297)
section.top_margin, section.bottom_margin = Mm(22), Mm(18)
section.left_margin = section.right_margin = Mm(18)
section.header_distance = section.footer_distance = Mm(8)
normal = doc.styles["Normal"]
normal.font.name, normal.font.size = "Arial", Pt(10)
normal.font.color.rgb = RGBColor.from_string("181818")
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.05
for name, size in [("Title", 24), ("Heading 1", 13), ("Heading 2", 11)]:
    style = doc.styles[name]
    style.font.name, style.font.size = "Arial", Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(PURPLE)
    style.paragraph_format.space_before = Pt(9)
    style.paragraph_format.space_after = Pt(5)
    style.paragraph_format.keep_with_next = True

# Explicit colours prevent a user's Office theme supplying a blue title rule.
title_border = OxmlElement("w:pBdr")
title_bottom = OxmlElement("w:bottom")
for key, value in {"val": "single", "sz": "8", "color": PINK, "space": "5"}.items():
    title_bottom.set(qn("w:" + key), value)
title_border.append(title_bottom)
title_props = doc.styles["Title"].element.get_or_add_pPr()
for old_border in list(title_props.findall(qn("w:pBdr"))):
    title_props.remove(old_border)
title_props.append(title_border)

header = section.header.paragraphs[0]
header.add_run().add_picture(str(ROOT / "src/assets/rws-logo.png"), width=Mm(32))
header.add_run("    DOCUMENT COMPARE").font.color.rgb = RGBColor.from_string(PURPLE)
header.runs[-1].font.size = Pt(9)
header.runs[-1].bold = True
header.paragraph_format.space_after = Pt(0)
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run("RWS  |  User quick guide  |  v1.0.21  •  ").font.size = Pt(8)
for instruction in ["PAGE", "NUMPAGES"]:
    if instruction == "NUMPAGES":
        footer.add_run(" / ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), instruction)
    footer._p.append(field)
border = OxmlElement("w:pBdr")
top = OxmlElement("w:top")
for key, value in {"val": "single", "sz": "10", "color": PINK, "space": "6"}.items():
    top.set(qn("w:" + key), value)
border.append(top)
footer._p.get_or_add_pPr().append(border)

def para(text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead:
        p.add_run(bold_lead).bold = True
        p.add_run(text)
    else:
        p.add_run(text)
    return p

def step(number, lead, text):
    p = para(text, str(number) + "  " + lead + " ")
    p.paragraph_format.space_after = Pt(5)

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    for cell, width in zip(t.rows[0].cells, widths):
        cell.width = Mm(width)
    for cell, label in zip(t.rows[0].cells, headers):
        cell.text = label
        shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), PURPLE)
        cell._tc.get_or_add_tcPr().append(shd)
        for r in cell.paragraphs[0].runs:
            r.font.bold = True; r.font.color.rgb = RGBColor(255, 255, 255)
    for values in rows:
        cells = t.add_row().cells
        for cell, value, width in zip(cells, values, widths):
            cell.width = Mm(width); cell.text = value
            shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), "F6F6F6")
            cell._tc.get_or_add_tcPr().append(shd)
    for row in t.rows:
        pr = row._tr.get_or_add_trPr(); pr.append(OxmlElement("w:cantSplit"))
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.space_before = Pt(4)
                for r in p.runs: r.font.size = Pt(9)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t

doc.add_paragraph("Your comparison, clearly", "Title")
para("Compare archive and new documents locally in desktop Word. Insertions are yellow; deletions and Word-recognised moves remain tracked.")
doc.add_heading("Get started", level=1)
para("Open the desktop shortcut. Desktop Word must be installed and activated. Browse DOCX, DOC or PDF files from local, company OneDrive or shared-drive folders.")
doc.add_heading("Which option should I use?", level=1)
table(["Choose", "When to use it"], [
    ("Full Document", "Use this by default for two files. Compares all included text and returns one set of figures."),
    ("EP", "Use for a EuroFile multi-language job when you need Description, Full specification and Claims figures separately. Use prepared Word files with clear section headings; translated claims are excluded."),
    ("Batch comparison tab", "Compare every later file with every earlier file in a folder. Choose Full Document or EP using the same rules above."),
    ("Multi Archive tab", "Compare one new document against several archives to find the best match. Choose Full Document or EP using the same rules above.")
], [35, 139])
para("For EP PDFs or files converted from PDF, use Full Document: section detection can be unreliable. Prepare both inputs to the same scope, removing covers, translations and separate drawings as needed.")
doc.add_heading("Single comparison", level=1)
step(1, "Choose Full Document or EP.", "Check that both inputs contain the intended source language and sections.")
step(2, "Browse for Archive, then New.", "Or drag and drop / enter a full file path. Archive is the earlier version; New is the version to assess.")
step(3, "Set PDF scan options if needed.", "Use local OCR only for PDFs containing scanned pages; see page 2.")
step(4, "Choose where to save.", "Keep the default folder, use Desktop/Browse, or enter a path. Local, OneDrive and shared-drive destinations are supported.")
step(5, "Click Compare documents.", "When finished, open the Word comparison. Its name is [new name] vs [archive name].docx; a matching TXT summary is saved alongside it.")
doc.add_heading("Read the three results", level=1)
table(["Result", "Meaning"], [
    ("Crossover", "Share of new words retained from the archive: 100 × (new words − inserted words) ÷ new words. Displayed to one decimal place."),
    ("New words", "Total final word count of the new file, or the selected EP scope."),
    ("Inserted words", "Word count of text marked as insertions by this comparison. Existing yellow highlighting alone does not count.")
], [37, 137])
para("Example: 1,000 new words and 250 inserted words give 75.0% crossover. A deletion alone does not reduce crossover; a recognised move is retained.")
para("In Word, select Review → Display for Review → All Markup to inspect deletions and moves as well as yellow insertions. Check the output before relying on its figures.")

doc.add_page_break()
doc.add_paragraph("Batch comparisons & useful checks", "Title")
doc.add_heading("Compare every pair in a folder", level=1)
step(1, "Open Batch comparison.", "Paste a full local, OneDrive or shared-drive folder path, or use Browse, then Load files. Subfolders and Word lock files are ignored.")
step(2, "Check the order.", "Filename prefixes such as 01, 02, 03 set the initial order. Use Move up, Move down or Remove to adjust the list. Changing the folder requires loading its files again.")
step(3, "Choose Full Document or EP for the batch.", "For PDF-converted EP jobs, choose Full Document and prepare every input to the same scope. Tick PDF OCR for each scanned PDF.")
step(4, "Choose Save batch to, then Start batch.", "Each later file is New and is compared against every earlier Archive. Keep the selected source files unchanged while the batch runs.")
step(5, "Review the outputs.", "A new Batch folder holds numbered pair folders with Word comparisons and TXT summaries. Open batch-results.csv for the combined table: one row per Full Document pair, or three scope rows per EP pair.")
para("Four files = six comparisons: 2 vs 1, 3 vs 1, 4 vs 1, 3 vs 2, 4 vs 2, 4 vs 3. Comparisons run one at a time; larger batches take longer.")
doc.add_heading("Language checks", level=1)
para("A known language mismatch stops a single comparison without output. A pop-up shows the scope and archive/new languages: select same-language inputs. Batch mismatch pairs fail while other pairs continue; one summary appears when finished or cancelled. Unknown/Mixed still stops for manual checking; in EP, check Description and original-language Claims separately.")
para("After confirming matching languages, use Language checked — compare for a single job, or Language checked — continue for the paused batch pair. Batch confirmation applies only to that pair. Skip pair continues without its comparison. A known mismatch cannot be overridden.")
doc.add_heading("PDFs, scans and ABBYY", level=1)
para("For selectable-text PDFs, leave OCR off. Scan mode reconstructs page sizes, text sizes, indents and spacing where possible. Review against the PDF: fonts, complex columns, tables and graphics may differ. Pages with no recognised words block an incomplete result.")
para("ABBYY can be used separately to prepare a checked, editable DOCX before comparison. The app does not launch ABBYY; its built-in PDF OCR uses Windows. PDF OCR does not recover text from images embedded in Word files.")
doc.add_heading("Save safely; recover when needed", level=1)
para("Processing and OCR use local working copies; originals are not overwritten. Existing source revisions are accepted on copies. Browse or enter UNC (\\\\server\\share\\folder) / mapped-drive (S:\\folder) paths for shared files and results. Keep the drive connected and check read/write permissions. OneDrive outputs may sync; follow company storage rules. Word’s connected features follow company settings.")
para("Cancel batch keeps completed outputs and records unfinished pairs. Failed or unfinished rows have blank figures, not zero. If a CSV cannot be saved, close it in Excel and choose Retry saving results; after stopping, use Save results table.")
para("Use Reset on either tab to start another job without closing the app. It clears selected files, on-screen results and language confirmation. Your job type, save folder and saved outputs stay. Reset is disabled during processing or a paused batch.")
para("If crossover is Unavailable, read the TXT/CSV notes and check the comparison. Empty new text or a genuine final-text mismatch has no usable percentage. For PDF-converted EP files, prefer Full Document; for native Word EP files, check Description and original-language Claims headings.")

doc.add_page_break()
doc.add_paragraph("Find the best archive", "Title")
para("Use Multi Archive when you have one new document and several possible archives. Each archive is compared with that same new document; archives are not compared with each other.")

doc.add_heading("Run a Multi Archive search", level=1)
step(1, "Open Multi Archive.", "Choose Full Document by default, or EP when separate section figures are needed and the Word headings are clear.")
step(2, "Add your archives.", "Use Add archives to select multiple files, or drop files into the list. You can also browse or enter a folder and choose Load files. Folder loading replaces the archive list; subfolders are ignored. Add archives adds files to the list and ignores duplicates.")
step(3, "Choose the New document below.", "Use Browse, drag and drop, or enter its full path. Select local OCR only if the new PDF contains scanned pages. Remove it from the archive list if it was loaded from the same folder.")
step(4, "Check the selection and save folder.", "Remove unwanted archives and select PDF OCR for scanned archive PDFs. The new file must be different from every archive. Local, company OneDrive and accessible shared drives are supported.")
step(5, "Click Find best match.", "The app compares each archive in turn, keeping a Word comparison and TXT summary for every successful pair.")

doc.add_heading("Read the ranked results", level=1)
table(["Result", "How to use it"], [
    ("Rank", "Highest verified crossover comes first. Exact score ties share a rank. Ranking uses unrounded percentages; display uses one decimal place. Rank so far is provisional during processing."),
    ("Full Document", "One row per archive with crossover, total new words and highlighted insertion words."),
    ("EP", "Three rows per archive. Ranking uses Full specification crossover; Description and Claims figures are also shown."),
    ("Best match", "The final summary names the best ranked archive, including ties. Open its comparison and check the markup before choosing it."),
    ("Unranked", "Failed, unfinished, empty-new-text or unavailable-crossover results have no rank. Available word counts alone do not establish a verified match.")
], [37, 137])
para("The new Multi Archive results folder contains numbered comparison folders and multi-archive-results.csv. The CSV follows the same ranking as the app and retains all three EP scope rows. Each Word comparison is named [new name] vs [archive name].docx.")

doc.add_heading("Check, cancel or start again", level=1)
para("A known language mismatch fails that archive while the remaining comparisons continue; one pop-up summarises mismatches at the end. Unknown/Mixed pauses for a manual check of the current pair. After confirming matching source languages, choose Language checked — continue, or use Skip archive. A known mismatch cannot be overridden.")
para("Cancel search keeps completed comparisons and ranks only those with verified figures; its best match is among the completed comparisons. Reset clears the new file, archive list and on-screen results while keeping the job type, save folder and saved files. Only one comparison tab can run at a time.")

doc.core_properties.title = "Document Compare — User quick guide"
doc.core_properties.subject = "Single, batch and Multi Archive comparisons, results and local document checks"
doc.core_properties.author = "RWS"
doc.core_properties.keywords = "RWS, Document Compare, quick guide, Full Document, WF, EP"
doc.save(OUT)
print(OUT)
